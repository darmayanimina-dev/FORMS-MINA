"""
FORMS - FAB OPEX Reconciliation & Monitoring System
Excel Generator: Standardized, accounting-grade Excel output via openpyxl.
Target file: Detail Realisasi OPEX - [bulan apa].xlsx
"""

import io
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_detail_excel(df_detail: pd.DataFrame, df_rekap: pd.DataFrame, period: str) -> io.BytesIO:
    """
    Generates a beautifully styled Excel workbook according to FAB financial standards.
    Sheet 1: Detail Transaksi (Row 1 Header: Batch Name, Trans. Date, Code Account, Account Name, Saldo, Description)
    Sheet 2: Rekapitulasi Akun (Code Account, Account Name, Total Saldo, % Proporsi)
    """
    wb = Workbook()
    
    # ----------------------------------------------------
    # Sheet 1: Detail Realisasi OPEX
    # ----------------------------------------------------
    ws_detail = wb.active
    ws_detail.title = "Detail Transaksi"
    ws_detail.views.sheetView[0].showGridLines = True

    # Color Palette & Styles
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Deep Corporate Navy
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10, color="0F172A")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid") # Light Slate
    negative_font = Font(name="Calibri", size=10, color="DC2626") # Red for minus values
    
    total_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    total_font = Font(name="Calibri", size=11, bold=True, color="0F172A")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    double_bottom_border = Border(
        left=thin_border_side, 
        right=thin_border_side, 
        top=Side(border_style="thin", color="0F172A"), 
        bottom=Side(border_style="double", color="0F172A")
    )

    # Headers strictly on Row 1 as requested:
    # 1. Batch Name, 2. Trans. Date, 3. Code Account, 4. Account Name, 5. Saldo, 6. Description
    headers = [
        "Batch Name",
        "Trans. Date",
        "Code Account",
        "Account Name",
        "Saldo",
        "Description"
    ]

    ws_detail.row_dimensions[1].height = 28
    for col_idx, header in enumerate(headers, 1):
        cell = ws_detail.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border

    # Data Rows
    current_row = 2
    for _, row_data in df_detail.iterrows():
        ws_detail.row_dimensions[current_row].height = 20
        is_zebra = (current_row % 2 == 0)
        row_fill = zebra_fill if is_zebra else None

        # 1. Batch Name
        c1 = ws_detail.cell(row=current_row, column=1, value=str(row_data.get("Batch Name", "") or ""))
        c1.alignment = Alignment(horizontal="center", vertical="center")
        
        # 2. Trans. Date
        c2 = ws_detail.cell(row=current_row, column=2, value=str(row_data.get("Trans. Date", "") or ""))
        c2.alignment = Alignment(horizontal="center", vertical="center")

        # 3. Code Account (Text formatted to prevent loss of leading zeros/special characters)
        c3 = ws_detail.cell(row=current_row, column=3, value=str(row_data.get("Code Account", "") or ""))
        c3.number_format = "@"
        c3.alignment = Alignment(horizontal="center", vertical="center")

        # 4. Account Name
        c4 = ws_detail.cell(row=current_row, column=4, value=str(row_data.get("Account Name", "") or ""))
        c4.alignment = Alignment(horizontal="left", vertical="center")

        # 5. Saldo (Standard accounting format)
        saldo_val = float(row_data.get("Saldo", 0.0) or 0.0)
        c5 = ws_detail.cell(row=current_row, column=5, value=saldo_val)
        c5.number_format = '_(* #,##0.00_);_(* (#,##0.00);_(* "-"??_);_(@_)'
        c5.alignment = Alignment(horizontal="right", vertical="center")
        if saldo_val < 0:
            c5.font = negative_font
        else:
            c5.font = data_font

        # 6. Description
        c6 = ws_detail.cell(row=current_row, column=6, value=str(row_data.get("Description", "") or ""))
        c6.alignment = Alignment(horizontal="left", vertical="center")

        for col_idx in range(1, 7):
            cell = ws_detail.cell(row=current_row, column=col_idx)
            if col_idx != 5:
                cell.font = data_font
            if row_fill:
                cell.fill = row_fill
            cell.border = cell_border

        current_row += 1

    # Total Summary Row
    ws_detail.row_dimensions[current_row].height = 24
    ws_detail.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
    total_label_cell = ws_detail.cell(row=current_row, column=1, value="TOTAL BEBAN OPERASIONAL (OPEX)")
    total_label_cell.font = total_font
    total_label_cell.fill = total_fill
    total_label_cell.alignment = Alignment(horizontal="right", vertical="center")
    total_label_cell.border = double_bottom_border

    for c in range(2, 5):
        ws_detail.cell(row=current_row, column=c).border = double_bottom_border
        ws_detail.cell(row=current_row, column=c).fill = total_fill

    # Total Formula
    total_val_cell = ws_detail.cell(row=current_row, column=5)
    if len(df_detail) > 0:
        total_val_cell.value = f"=SUM(E2:E{current_row-1})"
    else:
        total_val_cell.value = 0.0
    total_val_cell.number_format = '_(* #,##0.00_);_(* (#,##0.00);_(* "-"??_);_(@_)'
    total_val_cell.font = total_font
    total_val_cell.fill = total_fill
    total_val_cell.alignment = Alignment(horizontal="right", vertical="center")
    total_val_cell.border = double_bottom_border

    blank_end_cell = ws_detail.cell(row=current_row, column=6, value="")
    blank_end_cell.fill = total_fill
    blank_end_cell.border = double_bottom_border

    # Auto-fit Column Widths
    for col in ws_detail.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if cell.number_format and ('#,##0' in cell.number_format) and isinstance(cell.value, (int, float)):
                val_str = f"{cell.value:,.2f}"
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws_detail.column_dimensions[col_letter].width = max(max_len + 4, 14)
    ws_detail.column_dimensions['D'].width = max(ws_detail.column_dimensions['D'].width, 30)
    ws_detail.column_dimensions['F'].width = max(ws_detail.column_dimensions['F'].width, 35)

    # ----------------------------------------------------
    # Sheet 2: Rekapitulasi Pos Akun (Bonus Tim FAB)
    # ----------------------------------------------------
    ws_rekap = wb.create_sheet(title="Rekapitulasi Akun")
    ws_rekap.views.sheetView[0].showGridLines = True

    rekap_headers = ["No", "Code Account", "Account Name", "Total Saldo", "% Proporsi"]
    ws_rekap.row_dimensions[1].height = 28
    for col_idx, h in enumerate(rekap_headers, 1):
        c = ws_rekap.cell(row=1, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = cell_border

    r_row = 2
    sum_saldo = float(df_rekap["Total Saldo"].sum()) if not df_rekap.empty else 0.0

    for idx, r_data in df_rekap.iterrows():
        ws_rekap.row_dimensions[r_row].height = 20
        is_zebra = (r_row % 2 == 0)
        row_fill = zebra_fill if is_zebra else None
        
        saldo_val = float(r_data.get("Total Saldo", 0.0) or 0.0)
        pct = (saldo_val / sum_saldo * 100) if sum_saldo != 0 else 0.0

        # No
        ws_rekap.cell(row=r_row, column=1, value=idx + 1).alignment = Alignment(horizontal="center", vertical="center")
        
        # Code Account
        c_code = ws_rekap.cell(row=r_row, column=2, value=str(r_data.get("Code Account", "") or ""))
        c_code.number_format = "@"
        c_code.alignment = Alignment(horizontal="center", vertical="center")
        
        # Account Name
        ws_rekap.cell(row=r_row, column=3, value=str(r_data.get("Account Name", "") or "")).alignment = Alignment(horizontal="left", vertical="center")
        
        # Total Saldo
        c_saldo = ws_rekap.cell(row=r_row, column=4, value=saldo_val)
        c_saldo.number_format = '_(* #,##0.00_);_(* (#,##0.00);_(* "-"??_);_(@_)'
        c_saldo.alignment = Alignment(horizontal="right", vertical="center")

        # % Proporsi
        c_pct = ws_rekap.cell(row=r_row, column=5, value=pct / 100)
        c_pct.number_format = '0.00%'
        c_pct.alignment = Alignment(horizontal="right", vertical="center")

        for col_idx in range(1, 6):
            cell = ws_rekap.cell(row=r_row, column=col_idx)
            cell.font = data_font
            if row_fill:
                cell.fill = row_fill
            cell.border = cell_border

        r_row += 1

    # Rekap Total
    ws_rekap.row_dimensions[r_row].height = 24
    ws_rekap.merge_cells(start_row=r_row, start_column=1, end_row=r_row, end_column=3)
    ws_rekap.cell(row=r_row, column=1, value="TOTAL SEMUA").alignment = Alignment(horizontal="right", vertical="center")
    
    ws_rekap.cell(row=r_row, column=4, value=f"=SUM(D2:D{r_row-1})" if len(df_rekap) > 0 else 0.0)
    ws_rekap.cell(row=r_row, column=4).number_format = '_(* #,##0.00_);_(* (#,##0.00);_(* "-"??_);_(@_)'
    ws_rekap.cell(row=r_row, column=4).alignment = Alignment(horizontal="right", vertical="center")

    ws_rekap.cell(row=r_row, column=5, value=1.0)
    ws_rekap.cell(row=r_row, column=5).number_format = '0.00%'
    ws_rekap.cell(row=r_row, column=5).alignment = Alignment(horizontal="right", vertical="center")

    for col_idx in range(1, 6):
        cell = ws_rekap.cell(row=r_row, column=col_idx)
        cell.font = total_font
        cell.fill = total_fill
        cell.border = double_bottom_border

    # Auto-fit sheet 2
    for col in ws_rekap.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws_rekap.column_dimensions[col_letter].width = max(max_len + 4, 12)
    ws_rekap.column_dimensions['C'].width = max(ws_rekap.column_dimensions['C'].width, 32)

    # Save to BytesIO
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
